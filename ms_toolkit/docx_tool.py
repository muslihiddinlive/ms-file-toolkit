"""DOCX (Word) fayllarni o'qish va tahlil qilish funksiyalari."""

from docx import Document


def read_docx_text(file_path: str) -> dict:
    """Word hujjatidagi barcha matnni paragraflar bo'yicha o'qiydi."""
    doc = Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return {
        "file_path": file_path,
        "paragraph_count": len(paragraphs),
        "text": "\n".join(paragraphs),
    }


def read_docx_tables(file_path: str) -> dict:
    """Word hujjatidagi barcha jadvallarni o'qiydi."""
    doc = Document(file_path)
    tables = []
    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])
        tables.append(rows)
    return {"file_path": file_path, "table_count": len(tables), "tables": tables}


def get_docx_metadata(file_path: str) -> dict:
    """Word hujjatining metama'lumotlarini (muallif, sarlavha va h.k.) qaytaradi."""
    doc = Document(file_path)
    props = doc.core_properties
    return {
        "file_path": file_path,
        "title": props.title,
        "author": props.author,
        "created": str(props.created) if props.created else None,
        "modified": str(props.modified) if props.modified else None,
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
    }
