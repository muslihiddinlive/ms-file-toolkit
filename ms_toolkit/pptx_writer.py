"""PPTX (PowerPoint) fayllarni yaratish va tahrirlash (yozish) funksiyalari."""

import os
from pptx import Presentation
from pptx.util import Inches, Emu


def create_pptx(file_path: str, title: str = None, subtitle: str = None, overwrite: bool = False) -> dict:
    """Yangi PowerPoint (.pptx) taqdimot yaratadi (title-slide bilan boshlanadi)."""
    if os.path.exists(file_path) and not overwrite:
        return {"error": f"Fayl allaqachon mavjud: {file_path} (overwrite=True qiling)"}

    prs = Presentation()
    layout = prs.slide_layouts[0]  # Title Slide
    slide = prs.slides.add_slide(layout)
    if title:
        slide.shapes.title.text = title
    if subtitle and len(slide.placeholders) > 1:
        slide.placeholders[1].text = subtitle

    os.makedirs(os.path.dirname(os.path.abspath(file_path)) or ".", exist_ok=True)
    prs.save(file_path)

    return {"file_path": file_path, "status": "created", "slide_count": len(prs.slides)}


def add_pptx_slide(
    file_path: str,
    title: str = None,
    bullet_points: list = None,
    notes: str = None,
) -> dict:
    """Mavjud taqdimotga sarlavha + bullet-point matnli yangi slayd qo'shadi."""
    if not os.path.exists(file_path):
        return {"error": f"Fayl topilmadi: {file_path}"}

    prs = Presentation(file_path)
    layout = prs.slide_layouts[1]  # Title and Content
    slide = prs.slides.add_slide(layout)

    if title and slide.shapes.title:
        slide.shapes.title.text = title

    if bullet_points:
        body = slide.placeholders[1]
        tf = body.text_frame
        tf.text = bullet_points[0]
        for point in bullet_points[1:]:
            p = tf.add_paragraph()
            p.text = point

    if notes:
        slide.notes_slide.notes_text_frame.text = notes

    prs.save(file_path)
    return {"file_path": file_path, "status": "slide_added", "slide_count": len(prs.slides)}


def add_pptx_image(
    file_path: str,
    image_path: str,
    slide_index: int = None,
    left_inches: float = 1.0,
    top_inches: float = 1.5,
    width_inches: float = None,
) -> dict:
    """Taqdimotga rasm qo'shadi.

    slide_index berilmasa, oxirgi slaydga qo'shiladi. width_inches berilmasa,
    rasm o'z asl o'lchamida (piksellardan EMU'ga o'girib) qo'yiladi.
    """
    if not os.path.exists(file_path):
        return {"error": f"Taqdimot topilmadi: {file_path}"}
    if not os.path.exists(image_path):
        return {"error": f"Rasm fayli topilmadi: {image_path}"}

    prs = Presentation(file_path)
    if not prs.slides:
        return {"error": "Taqdimotda hech qanday slayd yo'q"}

    idx = slide_index if slide_index is not None else len(prs.slides) - 1
    if not 0 <= idx < len(prs.slides):
        return {"error": f"Noto'g'ri slide_index: {idx} (jami slaydlar: {len(prs.slides)})"}

    slide = prs.slides[idx]
    kwargs = {"width": Inches(width_inches)} if width_inches else {}
    slide.shapes.add_picture(image_path, Inches(left_inches), Inches(top_inches), **kwargs)

    prs.save(file_path)
    return {"file_path": file_path, "status": "image_added", "slide_index": idx}


def set_pptx_background_color(file_path: str, slide_index: int, hex_color: str) -> dict:
    """Belgilangan slaydning fon rangini o'rnatadi (hex_color: masalan 'FFFFFF')."""
    from pptx.dml.color import RGBColor

    if not os.path.exists(file_path):
        return {"error": f"Taqdimot topilmadi: {file_path}"}

    prs = Presentation(file_path)
    if not 0 <= slide_index < len(prs.slides):
        return {"error": f"Noto'g'ri slide_index: {slide_index} (jami slaydlar: {len(prs.slides)})"}

    slide = prs.slides[slide_index]
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(hex_color.lstrip("#").upper())

    prs.save(file_path)
    return {"file_path": file_path, "status": "background_set", "slide_index": slide_index}


def add_pptx_chart(
    file_path: str,
    chart_type: str,
    categories: list,
    series: dict,
    slide_index: int = None,
    title: str = None,
    left_inches: float = 1.0,
    top_inches: float = 1.5,
    width_inches: float = 8.0,
    height_inches: float = 4.5,
) -> dict:
    """Taqdimotga diagramma (bar/line/pie) qo'shadi.

    categories: X o'qi label'lari, masalan ["Yan", "Fev", "Mar"].
    series: {"Seriya nomi": [qiymatlar ro'yxati], ...} — bir yoki bir nechta seriya.
    slide_index berilmasa, oxirgi slaydga qo'shiladi.
    """
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    type_map = {
        "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "line": XL_CHART_TYPE.LINE,
        "pie": XL_CHART_TYPE.PIE,
    }
    xl_type = type_map.get(chart_type)
    if xl_type is None:
        return {"error": f"Noto'g'ri chart_type: {chart_type} (ruxsat etilgan: bar, line, pie)"}

    if not os.path.exists(file_path):
        return {"error": f"Taqdimot topilmadi: {file_path}"}

    prs = Presentation(file_path)
    if not prs.slides:
        return {"error": "Taqdimotda hech qanday slayd yo'q"}

    idx = slide_index if slide_index is not None else len(prs.slides) - 1
    if not 0 <= idx < len(prs.slides):
        return {"error": f"Noto'g'ri slide_index: {idx} (jami slaydlar: {len(prs.slides)})"}

    chart_data = CategoryChartData()
    chart_data.categories = categories
    for series_name, values in series.items():
        chart_data.add_series(series_name, values)

    slide = prs.slides[idx]
    graphic_frame = slide.shapes.add_chart(
        xl_type,
        Inches(left_inches), Inches(top_inches),
        Inches(width_inches), Inches(height_inches),
        chart_data,
    )

    if title:
        graphic_frame.chart.has_title = True
        graphic_frame.chart.chart_title.text_frame.text = title

    prs.save(file_path)
    return {"file_path": file_path, "status": "chart_added", "chart_type": chart_type, "slide_index": idx}
