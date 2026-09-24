"""Fayl formatlari o'rtasida konvertatsiya funksiyalari.

DOCX/PPTX -> PDF konvertatsiyasi `legacy_tool.py` bilan bir xil usulda
LibreOffice (`soffice`) orqali amalga oshiriladi — tizimda o'rnatilgan
bo'lishi shart: `apt-get install libreoffice`.

PDF -> rasm konvertatsiyasi **PyMuPDF** (`pymupdf`) orqali, tashqi dastursiz
(system binary shart emas) amalga oshiriladi: `pip install pymupdf`.
"""

import os
import subprocess
import tempfile
import uuid

import pymupdf


def _soffice_to_pdf(file_path: str, timeout: int = 60) -> str:
    """`file_path`ni (docx/pptx/xlsx/odt/...) vaqtinchalik papkada PDF'ga o'giradi.

    Qaytaradi: hosil bo'lgan PDF faylning to'liq yo'li. Chaqiruvchi vaqtinchalik
    faylni o'zi tozalashi kerak.
    """
    out_dir = os.path.join(tempfile.gettempdir(), f"ms_toolkit_convert_{uuid.uuid4().hex}")
    os.makedirs(out_dir, exist_ok=True)

    result = subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", out_dir, file_path],
        capture_output=True,
        text=True,
        timeout=timeout,
    )
    if result.returncode != 0:
        raise RuntimeError(f"LibreOffice konvertatsiya xatosi: {result.stderr or result.stdout}")

    base_name = os.path.splitext(os.path.basename(file_path))[0]
    pdf_path = os.path.join(out_dir, f"{base_name}.pdf")
    if not os.path.exists(pdf_path):
        raise RuntimeError(f"Konvertatsiya qilingan PDF topilmadi: {pdf_path}")

    return pdf_path


def convert_docx_to_pdf(file_path: str, output_path: str, overwrite: bool = False) -> dict:
    """Word (.docx) hujjatini PDF'ga o'giradi (LibreOffice orqali)."""
    return _convert_office_to_pdf(file_path, output_path, overwrite)


def convert_pptx_to_pdf(file_path: str, output_path: str, overwrite: bool = False) -> dict:
    """PowerPoint (.pptx) taqdimotini PDF'ga o'giradi (LibreOffice orqali)."""
    return _convert_office_to_pdf(file_path, output_path, overwrite)


def convert_xlsx_to_pdf(file_path: str, output_path: str, overwrite: bool = False) -> dict:
    """Excel (.xlsx) faylini PDF'ga o'giradi (LibreOffice orqali)."""
    return _convert_office_to_pdf(file_path, output_path, overwrite)


def _convert_office_to_pdf(file_path: str, output_path: str, overwrite: bool) -> dict:
    if not os.path.exists(file_path):
        return {"error": f"Fayl topilmadi: {file_path}"}
    if os.path.exists(output_path) and not overwrite:
        return {"error": f"Fayl allaqachon mavjud: {output_path} (overwrite=True qiling)"}

    tmp_pdf = None
    try:
        tmp_pdf = _soffice_to_pdf(file_path)
        os.makedirs(os.path.dirname(os.path.abspath(output_path)) or ".", exist_ok=True)
        os.replace(tmp_pdf, output_path)
        tmp_pdf = None
        return {"file_path": output_path, "status": "converted", "source_file": file_path}
    except subprocess.TimeoutExpired:
        return {"error": "Konvertatsiya vaqti tugadi (timeout)"}
    except Exception as e:
        return {"error": f"{type(e).__name__}: {e}"}
    finally:
        if tmp_pdf and os.path.exists(tmp_pdf):
            out_dir = os.path.dirname(tmp_pdf)
            try:
                os.remove(tmp_pdf)
                os.rmdir(out_dir)
            except OSError:
                pass


def convert_pdf_to_images(
    file_path: str,
    output_dir: str,
    image_format: str = "png",
    dpi: int = 150,
) -> dict:
    """PDF faylning har bir sahifasini alohida rasm faylga aylantiradi (PyMuPDF)."""
    if not os.path.exists(file_path):
        return {"error": f"Fayl topilmadi: {file_path}"}

    os.makedirs(output_dir, exist_ok=True)
    output_files = []

    try:
        doc = pymupdf.open(file_path)
        for i, page in enumerate(doc, start=1):
            pix = page.get_pixmap(dpi=dpi)
            out_path = os.path.join(output_dir, f"page_{i}.{image_format}")
            pix.save(out_path)
            output_files.append(out_path)
        doc.close()
    except Exception as e:
        return {"error": f"{type(e).__name__}: {e}"}

    return {"file_path": file_path, "status": "converted", "output_files": output_files, "page_count": len(output_files)}


def convert_pptx_to_images(
    file_path: str,
    output_dir: str,
    image_format: str = "png",
    dpi: int = 150,
) -> dict:
    """PowerPoint taqdimotining har bir slaydini rasmga aylantiradi (avval PDF'ga, keyin rasmga)."""
    if not os.path.exists(file_path):
        return {"error": f"Fayl topilmadi: {file_path}"}

    tmp_pdf = None
    try:
        tmp_pdf = _soffice_to_pdf(file_path)
        result = convert_pdf_to_images(tmp_pdf, output_dir, image_format=image_format, dpi=dpi)
        result["file_path"] = file_path
        return result
    except subprocess.TimeoutExpired:
        return {"error": "Konvertatsiya vaqti tugadi (timeout)"}
    except Exception as e:
        return {"error": f"{type(e).__name__}: {e}"}
    finally:
        if tmp_pdf and os.path.exists(tmp_pdf):
            out_dir = os.path.dirname(tmp_pdf)
            try:
                os.remove(tmp_pdf)
                os.rmdir(out_dir)
            except OSError:
                pass


def convert_xlsx_to_csv(file_path: str, output_dir: str, delimiter: str = ",") -> dict:
    """Excel (.xlsx) faylning har bir varag'ini alohida CSV faylga eksport qiladi."""
    if not os.path.exists(file_path):
        return {"error": f"Fayl topilmadi: {file_path}"}

    import csv
    import openpyxl

    os.makedirs(output_dir, exist_ok=True)
    output_files = []

    wb = openpyxl.load_workbook(file_path, data_only=True)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        safe_name = "".join(c if c.isalnum() or c in " _-" else "_" for c in sheet_name)
        out_path = os.path.join(output_dir, f"{safe_name}.csv")
        with open(out_path, "w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f, delimiter=delimiter)
            for row in ws.iter_rows(values_only=True):
                writer.writerow(["" if v is None else v for v in row])
        output_files.append(out_path)

    return {"file_path": file_path, "status": "converted", "output_files": output_files, "sheet_count": len(output_files)}
