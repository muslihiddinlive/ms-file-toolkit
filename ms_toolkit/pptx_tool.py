"""PPTX (PowerPoint) fayllarni o'qish va tahlil qilish funksiyalari."""

from pptx import Presentation


def read_pptx_text(file_path: str) -> dict:
    """PowerPoint taqdimotidagi har bir slayddan matnni ajratib oladi."""
    prs = Presentation(file_path)
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs)
                    if text.strip():
                        texts.append(text)
        slides.append({"slide_number": i, "text": "\n".join(texts)})
    return {"file_path": file_path, "slide_count": len(slides), "slides": slides}


def get_pptx_slide_count(file_path: str) -> dict:
    """Taqdimotdagi slaydlar sonini qaytaradi."""
    prs = Presentation(file_path)
    return {"file_path": file_path, "slide_count": len(prs.slides)}


def extract_pptx_notes(file_path: str) -> dict:
    """Har bir slaydning speaker notes (izoh) matnini ajratib oladi."""
    prs = Presentation(file_path)
    notes = []
    for i, slide in enumerate(prs.slides, start=1):
        if slide.has_notes_slide:
            note_text = slide.notes_slide.notes_text_frame.text
        else:
            note_text = ""
        notes.append({"slide_number": i, "notes": note_text})
    return {"file_path": file_path, "notes": notes}
