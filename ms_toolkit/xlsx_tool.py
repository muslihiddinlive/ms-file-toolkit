"""XLSX (Excel) fayllarni o'qish va tahlil qilish funksiyalari."""

from openpyxl import load_workbook


def get_xlsx_sheet_names(file_path: str) -> dict:
    """Excel faylidagi barcha varaq (sheet) nomlarini qaytaradi."""
    wb = load_workbook(file_path, read_only=True, data_only=True)
    return {"file_path": file_path, "sheet_names": wb.sheetnames}


def read_xlsx_data(file_path: str, sheet_name: str = None, max_rows: int = 100) -> dict:
    """Excel faylidan ma'lumotlarni jadval (list of dict) ko'rinishida o'qiydi."""
    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active

    rows_iter = ws.iter_rows(values_only=True)
    try:
        headers = next(rows_iter)
    except StopIteration:
        return {"file_path": file_path, "sheet_name": ws.title, "rows": []}

    data = []
    for i, row in enumerate(rows_iter):
        if i >= max_rows:
            break
        data.append(dict(zip(headers, row)))

    return {
        "file_path": file_path,
        "sheet_name": ws.title,
        "headers": list(headers),
        "row_count": len(data),
        "rows": data,
    }


def get_xlsx_summary(file_path: str) -> dict:
    """Excel faylining har bir varag'i uchun o'lcham (qator/ustun) xulosasini beradi."""
    wb = load_workbook(file_path, read_only=True, data_only=True)
    summary = []
    for name in wb.sheetnames:
        ws = wb[name]
        summary.append(
            {"sheet_name": name, "max_row": ws.max_row, "max_column": ws.max_column}
        )
    return {"file_path": file_path, "sheets": summary}
