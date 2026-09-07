"""XLSX (Excel) fayllarni yaratish va tahrirlash (yozish) funksiyalari."""

import os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


def create_xlsx(
    file_path: str,
    sheet_name: str = "Sheet1",
    headers: list = None,
    rows: list = None,
    overwrite: bool = False,
) -> dict:
    """Yangi Excel (.xlsx) fayl yaratadi.

    headers: ustun sarlavhalari ro'yxati, masalan ["Ism", "Yosh"].
    rows: har biri bitta qatorni ifodalovchi ro'yxatlar ro'yxati,
        masalan [["Ali", 25], ["Vali", 30]].
    """
    if os.path.exists(file_path) and not overwrite:
        return {"error": f"Fayl allaqachon mavjud: {file_path} (overwrite=True qiling)"}

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    if headers:
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)

    for row in rows or []:
        ws.append(row)

    os.makedirs(os.path.dirname(os.path.abspath(file_path)) or ".", exist_ok=True)
    wb.save(file_path)

    return {
        "file_path": file_path,
        "status": "created",
        "sheet_name": sheet_name,
        "row_count": len(rows or []),
    }


def add_xlsx_sheet(file_path: str, sheet_name: str, headers: list = None, rows: list = None) -> dict:
    """Mavjud Excel faylga yangi varaq (sheet) qo'shadi."""
    if not os.path.exists(file_path):
        return {"error": f"Fayl topilmadi: {file_path}"}

    wb = load_workbook(file_path)
    if sheet_name in wb.sheetnames:
        return {"error": f"'{sheet_name}' nomli varaq allaqachon mavjud"}

    ws = wb.create_sheet(sheet_name)
    if headers:
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
    for row in rows or []:
        ws.append(row)

    wb.save(file_path)
    return {"file_path": file_path, "status": "sheet_added", "sheet_name": sheet_name}


def append_xlsx_rows(file_path: str, rows: list, sheet_name: str = None) -> dict:
    """Mavjud Excel faylning bir varag'iga yangi qatorlar qo'shadi (oxiriga)."""
    if not os.path.exists(file_path):
        return {"error": f"Fayl topilmadi: {file_path}"}

    wb = load_workbook(file_path)
    ws = wb[sheet_name] if sheet_name else wb.active

    for row in rows or []:
        ws.append(row)

    wb.save(file_path)
    return {"file_path": file_path, "status": "appended", "sheet_name": ws.title, "rows_added": len(rows or [])}


def set_xlsx_formula(file_path: str, cell: str, formula: str, sheet_name: str = None) -> dict:
    """Belgilangan katakka Excel formulasi yozadi (masalan '=SUM(A1:A10)')."""
    if not os.path.exists(file_path):
        return {"error": f"Fayl topilmadi: {file_path}"}
    if not formula.startswith("="):
        formula = "=" + formula

    wb = load_workbook(file_path)
    ws = wb[sheet_name] if sheet_name else wb.active
    ws[cell] = formula
    wb.save(file_path)
    return {"file_path": file_path, "status": "formula_set", "cell": cell, "formula": formula, "sheet_name": ws.title}


def format_xlsx_cells(
    file_path: str,
    cell_range: str,
    bold: bool = None,
    bg_color: str = None,
    sheet_name: str = None,
) -> dict:
    """Katak diapazoniga formatlash qo'llaydi (qalin matn va/yoki fon rangi).

    cell_range: masalan 'A1' yoki 'A1:C1'.
    bg_color: 6 xonali hex rang kodi, masalan 'FFFF00' (sariq).
    """
    if not os.path.exists(file_path):
        return {"error": f"Fayl topilmadi: {file_path}"}

    wb = load_workbook(file_path)
    ws = wb[sheet_name] if sheet_name else wb.active

    fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid") if bg_color else None

    for row in ws[cell_range]:
        for cell in row:
            if bold is not None:
                cell.font = Font(bold=bold)
            if fill is not None:
                cell.fill = fill

    wb.save(file_path)
    return {"file_path": file_path, "status": "formatted", "cell_range": cell_range, "sheet_name": ws.title}


def freeze_xlsx_panes(file_path: str, cell: str = "A2", sheet_name: str = None) -> dict:
    """Excel varag'ida qatordagi/ustundagi sarlavhalarni muzlatadi (scroll qilganda ko'rinib turishi uchun).

    cell: muzlatish chegarasi, masalan 'A2' — 1-qator (sarlavha) muzlaydi.
    """
    if not os.path.exists(file_path):
        return {"error": f"Fayl topilmadi: {file_path}"}

    wb = load_workbook(file_path)
    ws = wb[sheet_name] if sheet_name else wb.active
    ws.freeze_panes = cell
    wb.save(file_path)
    return {"file_path": file_path, "status": "panes_frozen", "cell": cell, "sheet_name": ws.title}
